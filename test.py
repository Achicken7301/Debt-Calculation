from docx import Document
from docx.shared import Pt
import pandas as pd

# Tạo một tài liệu mới
doc = Document()

# Thêm tiêu đề cho tài liệu
doc.add_heading("Bảng Sản Phẩm", level=1)

# Thêm một bảng với 1 hàng tiêu đề và 4 cột
table = doc.add_table(rows=1, cols=4)

# Lấy hàng tiêu đề
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "Tên Sản Phẩm"
hdr_cells[1].text = "Đơn Giá"
hdr_cells[2].text = "Số Lượng"
hdr_cells[3].text = "Tổng"

# Định dạng tiêu đề bảng
for cell in hdr_cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            run.font.size = Pt(12)

# Thêm dữ liệu mẫu vào bảng
df = pd.read_csv("./sample file.csv")
print(df)
products = [tuple(row) for row in df.values]

# products = [
#     ("Sản Phẩm A", "100.000", "2", "200.000"),
#     ("Sản Phẩm B", "200.000", "1", "200.000"),
#     ("Sản Phẩm C", "150.000", "3", "450.000"),
# ]

for product in products:
    row_cells = table.add_row().cells
    row_cells[0].text = str(product[0])
    row_cells[1].text = str(product[1])
    row_cells[2].text = str(product[2])
    row_cells[3].text = str(product[3])

doc.save("bang_san_pham.docx")
